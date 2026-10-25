"""Asset validation, spreadsheet templates and local persistence."""
import ipaddress
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook

HEADERS = ["设备名称", "IP地址", "品牌", "SSH端口", "用户名", "密码", "分组"]
BRANDS = {"华为": "huawei", "huawei": "huawei", "华三": "hp_comware", "h3c": "hp_comware", "comware": "hp_comware", "hp_comware": "hp_comware"}

class AssetError(ValueError): pass

def expand_range(value, limit=4096):
    value = str(value).strip()
    if "/" in value:
        return expand_cidr(value, limit)
    parts = value.split("-")
    if len(parts) != 2:
        raise AssetError("请输入 CIDR、192.168.1.1-10 或完整起止 IPv4 地址")
    start, end = (part.strip() for part in parts)
    if end.isascii() and end.isdecimal():
        if not 0 <= int(end) <= 255:
            raise AssetError("简写结束地址必须为0-255")
        end = start.rsplit(".", 1)[0] + "." + end
    try:
        first, last = int(ipaddress.IPv4Address(start)), int(ipaddress.IPv4Address(end))
    except ipaddress.AddressValueError:
        raise AssetError("起止地址必须为有效 IPv4 地址")
    if first > last:
        raise AssetError("起始地址不能大于结束地址")
    count = last - first + 1
    if count > limit:
        raise AssetError(f"范围包含 {count} 个地址，单次上限 {limit}")
    return [str(ipaddress.IPv4Address(n)) for n in range(first, last + 1)]

def expand_cidr(value, limit=4096):
    try: net = ipaddress.ip_network(value.strip(), strict=False)
    except ValueError as e: raise AssetError(f"无效 IPv4 网段: {e}")
    if net.version != 4: raise AssetError("仅支持 IPv4 网段")
    count = net.num_addresses if net.prefixlen >= 31 else max(0, net.num_addresses - 2)
    if count > limit: raise AssetError(f"网段包含 {count} 个可用地址，单次上限 {limit}")
    return [str(x) for x in net.hosts()]

def normalize(record):
    r = dict(record)
    r["host"] = str(r.get("host", "")).strip()
    try: ipaddress.IPv4Address(r["host"])
    except ValueError: raise AssetError(f"无效 IPv4 地址: {r['host']}")
    vendor_value = str(r.get("vendor", "") or "").strip().lower()
    vendor = BRANDS.get(vendor_value)
    if vendor_value and not vendor: raise AssetError("品牌必须为华为或H3C")
    r["vendor"] = vendor or ""
    r["name"] = str(r.get("name") or r["host"]).strip()
    r["username"] = str(r.get("username", "")).strip()
    r["password"] = str(r.get("password", ""))
    if not r["username"] or not r["password"]: raise AssetError("用户名和密码不能为空")
    try: r["port"] = int(r.get("port", 22))
    except (TypeError, ValueError): raise AssetError("SSH端口必须为1-65535整数")
    if not 1 <= r["port"] <= 65535: raise AssetError("SSH端口必须为1-65535整数")
    r["group"] = str(r.get("group", "") or "未分组")
    return r

def write_template(path):
    wb = Workbook(); ws = wb.active; ws.title = "设备清单"; ws.append(HEADERS)
    ws.append(["核心交换机", "192.168.1.1", "华为", 22, "admin", "请替换为真实密码", "总部"])
    info = wb.create_sheet("说明"); info.append(["每行一台设备；密码明文保存，请妥善保护。品牌可留空，登录后识别华为或H3C。SSH端口列为连接端口：先尝试SSH，登录失败后在同端口尝试Telnet（明文传输）。"])
    wb.save(path)

def import_xlsx(path):
    wb = load_workbook(path, read_only=True, data_only=True); ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True)); wb.close()
    if not rows: raise AssetError("表格为空")
    header = {str(v).strip(): i for i, v in enumerate(rows[0]) if v is not None}
    mapping = {"name":"设备名称", "host":"IP地址", "vendor":"品牌", "port":"SSH端口", "username":"用户名", "password":"密码", "group":"分组"}
    missing = [v for v in mapping.values() if v not in header and v not in ("品牌", "分组")]
    if missing: raise AssetError("缺少表头: " + "、".join(missing))
    results=[]; errors=[]
    for line,row in enumerate(rows[1:], 2):
        if not row or all(v is None or str(v).strip()=="" for v in row): continue
        raw={k:(row[header[h]] if header.get(h) is not None and header[h] < len(row) else "") for k,h in mapping.items()}
        try: results.append(normalize(raw))
        except AssetError as e: errors.append(f"第{line}行: {e}")
    if errors: raise AssetError("导入校验失败：\n"+"\n".join(errors))
    if not results: raise AssetError("没有设备数据")
    return results

class AssetStore:
    def __init__(self, path): self.path=Path(path); self.assets=[]; self.history=[]; self.load()
    def load(self):
        if self.path.exists():
            data=json.loads(self.path.read_text(encoding="utf-8")); self.assets=data.get("assets", []); self.history=data.get("history", [])
    def save(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        tmp=self.path.with_suffix(".tmp"); tmp.write_text(json.dumps({"assets":self.assets,"history":self.history}, ensure_ascii=False, indent=2), encoding="utf-8"); tmp.replace(self.path)
